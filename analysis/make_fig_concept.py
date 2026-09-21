"""Figure 1: POPE responses -> hits and false alarms -> H, FA -> d', c -> stages.

Rebuilt 2026-09-21 from co-author review (Jae-Ho Lee):
  * a concrete POPE example, connecting "object present + Yes -> Hit" and
    "object absent + Yes -> False alarm", consistent with the main text;
  * panel (b) makes object present/absent, the criterion c, the false-alarm
    region and d' explicit, and shows two scenarios with the SAME accuracy but
    different response patterns;
  * the flow runs POPE responses -> Hit/False alarm -> H/FA -> d'/c, and panel
    (c) shows why d' and c are tracked separately over stages;
  * Times New Roman, sentence case, one set of panel labels, sizes and line
    widths, a tight layout, and an editable PowerPoint version.

Everything is REAL or COMPUTED, nothing is hand-placed data:
  (a) one POPE-adversarial image (COCO val2014 283412) on which the untuned
      LLaVA-1.5-7B base produces all four outcomes -- dog: Yes (hit), book: Yes
      (false alarm; the image has a newspaper), bed: No (miss), chair: No
      (correct rejection). Read from results_fs/llava15_base/pope_gen.jsonl.
  (b) two equal-variance Gaussian scenarios whose accuracies are asserted equal
      before anything is drawn. B's criterion is solved for by bisection.
  (c) the median-drift sequential UCIT run (seq|o3|s31; chosen by rank, not by
      look), change from the untuned base in SD units, from fs_aggregate.json.

ONE LAYOUT, TWO RENDERERS. The figure is a list of primitives in inch
coordinates (text runs, boxes, lines, arrows, filled polygons, an image). The
same list is drawn by matplotlib (the PDF the paper includes) and by
python-pptx (native, editable PowerPoint shapes), so the two cannot drift.
A layout check measures every rendered text box and fails the build on any
overlap or on text leaving its container.

    python3 analysis/make_fig_concept.py
      -> analysis/out/fig_concept.{pdf,png}, analysis/out/fig_concept.pptx
"""
import json
import math
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Polygon

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import fig_style as FS  # noqa: E402

FS.apply_style()
OUT = os.path.join(HERE, "out")
IMG = os.path.join(HERE, "assets", "pope_COCO_val2014_000000283412.jpg")
AGG = os.path.join(HERE, "readout", "fs_aggregate.json")

W, H = 5.50, 1.90                     # printed size: included at \textwidth
FONT = "Times New Roman"

INK, MUT, RULE = "#1A1A1A", "#5A6068", "#C9CED6"
PRES, ABS, ACC = "#0072B2", "#7D8590", "#D55E00"
PRES_FILL, FA_FILL, NEU_FILL, KEY_FILL = "#E4EEF7", "#FBE6DA", "#F3F4F6", "#FFF4D6"


# --------------------------------------------------------------------------
# numbers
# --------------------------------------------------------------------------
def Phi(z):
    return 0.5 * (1 + math.erf(z / math.sqrt(2)))


def pdf(x, mu):
    return math.exp(-0.5 * (x - mu) ** 2) / math.sqrt(2 * math.pi)


def rates(dprime, c):
    """Equal-variance SDT with means at -d'/2 (absent) and +d'/2 (present)."""
    m = dprime / 2
    return 1 - Phi(c - m), 1 - Phi(c + m)        # H, FA


def accuracy(dprime, c):
    H_, FA_ = rates(dprime, c)
    return 0.5 * (H_ + 1 - FA_)                   # POPE is balanced 4500/4500


A = dict(dprime=1.0, c=0.0)
B = dict(dprime=2.0, c=None)
lo, hi = -6.0, 0.0                                # liberal branch: c < 0
for _ in range(200):
    mid = 0.5 * (lo + hi)
    if accuracy(B["dprime"], mid) > accuracy(**A):
        hi = mid
    else:
        lo = mid
B["c"] = 0.5 * (lo + hi)
for S in (A, B):
    S["H"], S["FA"] = rates(S["dprime"], S["c"])
    S["acc"] = accuracy(S["dprime"], S["c"])
assert abs(A["acc"] - B["acc"]) < 5e-4, "panel (b) needs equal accuracies: %.4f vs %.4f" % (A["acc"], B["acc"])
assert B["FA"] > A["FA"] + 0.2, "B must hallucinate visibly more than A for the panel to make its point"

agg = json.load(open(AGG))["backbones"]["llava15"]
base = agg["base"]["pope"]
CELL = "seq|o3|s31"


def _drift(cell):
    cs = [cell[str(s)]["pope"]["c"] for s in range(1, 7)]
    return sum(abs(cs[i] - cs[i - 1]) for i in range(1, 6))


ranked = sorted((_drift(v), k) for k, v in agg["arms"].items() if k.startswith("seq|"))
assert ranked[len(ranked) // 2][1] == CELL, "panel (c) must show the MEDIAN-drift run; got %r" % ranked
run = agg["arms"][CELL]
DC = [0.0] + [run[str(s)]["pope"]["c"] - base["c"] for s in range(1, 7)]
DD = [0.0] + [run[str(s)]["pope"]["dprime"] - base["dprime"] for s in range(1, 7)]
ACCS = [0.5 * (base["H"] + 1 - base["FA"])] + [0.5 * (run[str(s)]["pope"]["H"] + 1 - run[str(s)]["pope"]["FA"])
                                                for s in range(1, 7)]


# --------------------------------------------------------------------------
# primitives (inches, origin bottom-left)
# --------------------------------------------------------------------------
P = []


def text(x, y, runs, size=7.5, ha="left", va="center", color=INK, rot=0, box=None, name=None):
    """runs: list of (string, style[, color]); style in '', 'i', 'b', 'bi'."""
    runs = [(r[0], r[1], r[2] if len(r) > 2 else color) for r in runs]
    P.append(dict(k="text", x=x, y=y, runs=runs, size=size, ha=ha, va=va, rot=rot, box=box, name=name))


def rect(x0, y0, x1, y1, fill=None, edge=None, lw=0.6, r=0.0):
    P.append(dict(k="rect", x0=x0, y0=y0, x1=x1, y1=y1, fill=fill, edge=edge, lw=lw, r=r))


def line(pts, color=INK, lw=0.6, dash=False):
    P.append(dict(k="line", pts=pts, color=color, lw=lw, dash=dash))


def arrow(p0, p1, color=INK, lw=0.7, both=False):
    P.append(dict(k="arrow", p0=p0, p1=p1, color=color, lw=lw, both=both))


def poly(pts, fill, alpha=1.0, edge=None, lw=0.0):
    P.append(dict(k="poly", pts=pts, fill=fill, alpha=alpha, edge=edge, lw=lw))


def dot(x, y, color, d=0.035):
    P.append(dict(k="dot", x=x, y=y, color=color, d=d))


def image(path, x0, y0, x1, y1):
    P.append(dict(k="image", path=path, x0=x0, y0=y0, x1=x1, y1=y1))


def frac(xc, y, num, den, size=7.5):
    """A stacked fraction built from native pieces, so it stays editable in PPT."""
    text(xc, y + 0.068, num, size=size, ha="center")
    text(xc, y - 0.068, den, size=size, ha="center")
    P.append(dict(k="fracbar", xc=xc, y=y, num=num, den=den, size=size))


# --------------------------------------------------------------------------
# layout  (every coordinate below is a printed inch; H chosen so the body stays
# inside 9 pages -- the old figure printed 1.56 in tall at 0.72\\textwidth)
# --------------------------------------------------------------------------
AX0, AX1 = 0.00, 1.92          # panel (a)
BX0, BX1 = 2.04, 4.33          # panel (b)
CX0, CX1 = 4.42, 5.50          # panel (c)
TITLE_Y = 1.835

# ---- (a) POPE responses -> hits and false alarms -> H and FA ---------------
text(AX0, TITLE_Y, [("(a)", "b"), (" POPE responses", "")], size=8.5, name="title_a")
image(IMG, 0.01, 1.17, 0.73, 1.65)
rect(0.01, 1.17, 0.73, 1.65, edge=RULE, lw=0.5)
qx = 0.81
text(qx, 1.595, [("POPE asks, e.g.", "")], size=7, color=MUT)
text(qx, 1.475, [("\u201cIs there a dog in", "")], size=7.5)
text(qx, 1.360, [("the image?\u201d", "")], size=7.5)
text(qx, 1.225, [("Half present, half absent", "")], size=7, color=MUT)

HDR_Y = 1.065
C1 = (0.30, 1.09)
C2 = (1.13, 1.92)
YES = (0.70, 0.97)
NO = (0.39, 0.66)
text(sum(C1) / 2, HDR_Y, [("Object present", "", PRES)], size=7.5, ha="center")
text(sum(C2) / 2, HDR_Y, [("Object absent", "", ABS)], size=7.5, ha="center")
text(0.0, HDR_Y, [("Answer", "")], size=7, ha="left", color=MUT)
text(0.14, sum(YES) / 2, [("Yes", "")], size=7.5, ha="center")
text(0.14, sum(NO) / 2, [("No", "")], size=7.5, ha="center")


def cell(cx, cy, fill, edge, head, head_style, head_color, example):
    rect(cx[0], cy[0], cx[1], cy[1], fill=fill, edge=edge, lw=0.7, r=0.035)
    mx = sum(cx) / 2
    box = (cx[0], cy[0], cx[1], cy[1])
    text(mx, cy[1] - 0.085, [(head, head_style, head_color)], size=7.5, ha="center", box=box)
    text(mx, cy[0] + 0.075, example, size=7, ha="center", color=MUT, box=box)


cell(C1, YES, PRES_FILL, PRES, "Hit", "b", PRES, [("dog: \u201cYes\u201d", "")])
cell(C2, YES, FA_FILL, ACC, "False alarm", "b", ACC, [("book: \u201cYes\u201d", "")])
cell(C1, NO, NEU_FILL, RULE, "Miss", "", INK, [("bed: \u201cNo\u201d", "")])
cell(C2, NO, NEU_FILL, RULE, "Correct rejection", "", INK, [("chair: \u201cNo\u201d", "")])

RATE_Y = 0.17
text(0.60, RATE_Y, [("H", "i"), (" =", "")], size=7.5, ha="right")
frac(0.83, RATE_Y, [("hits", "")], [("present", "")], size=7)
text(1.29, RATE_Y, [("FA =", "")], size=7.5, ha="right")
frac(1.61, RATE_Y, [("false alarms", "")], [("absent", "")], size=7)

# ---- (b) same accuracy, different responses --------------------------------
text(BX0, TITLE_Y, [("(b)", "b"), (" Same accuracy, different responses", "")], size=8.5, name="title_b")
text(BX0, 1.690, [("Discriminability ", ""), ("d", "i"), ("\u2032 = ", ""), ("z", "i"), ("(", ""), ("H", "i"),
                  (") \u2212 ", ""), ("z", "i"), ("(FA)", "")], size=7.5)
text(BX0, 1.575, [("Criterion ", ""), ("c", "i"), (" = \u2212\u00bd[", ""), ("z", "i"), ("(", ""), ("H", "i"),
                  (") + ", ""), ("z", "i"), ("(FA)]", "")], size=7.5)

STX = 3.55                      # right column: legend on top, then each scenario's numbers
P.append(dict(k="legend", x=STX, y=1.700, dy=-0.113, size=7.5,
              items=[(ABS, 0.45, "Object absent"), (PRES, 0.45, "Object present"), (ACC, 0.75, "False alarms")]))

PLX0, PLX1 = 2.18, 3.42         # plot strip; the scenario letter sits to its left
XR = (-4.0, 4.0)
YR = (-0.02, 0.60)


def plot_scenario(S, y0, y1, label):
    sx = lambda v: PLX0 + (v - XR[0]) / (XR[1] - XR[0]) * (PLX1 - PLX0)
    sy = lambda v: y0 + (v - YR[0]) / (YR[1] - YR[0]) * (y1 - y0)
    m = S["dprime"] / 2
    xs = [XR[0] + i * (XR[1] - XR[0]) / 240 for i in range(241)]
    text(BX0 + 0.04, (y0 + y1) / 2, [(label, "b")], size=8.5, ha="center")
    absent = [(sx(x), sy(pdf(x, -m))) for x in xs]
    present = [(sx(x), sy(pdf(x, +m))) for x in xs]
    poly([(sx(XR[0]), sy(0))] + absent + [(sx(XR[1]), sy(0))], ABS, alpha=0.22)
    poly([(sx(XR[0]), sy(0))] + present + [(sx(XR[1]), sy(0))], PRES, alpha=0.16)
    fa = [(sx(x), sy(pdf(x, -m))) for x in xs if x >= S["c"]]
    poly([(sx(S["c"]), sy(0)), (sx(S["c"]), sy(pdf(S["c"], -m)))] + fa + [(sx(XR[1]), sy(0))], ACC, alpha=0.55)
    line(absent, ABS, 1.1)
    line(present, PRES, 1.1)
    line([(sx(XR[0]), sy(0)), (sx(XR[1]), sy(0))], INK, 0.6)
    line([(sx(S["c"]), sy(0)), (sx(S["c"]), sy(0.425))], ACC, 1.3)
    text(sx(S["c"]), sy(0) - 0.062, [("c", "i", ACC)], size=8, ha="center")
    ya = sy(0.462)
    arrow((sx(-m), ya), (sx(+m), ya), INK, 0.7, both=True)
    text(sx(0), sy(0.545), [("d", "i"), ("\u2032", "")], size=8, ha="center")
    bt, bb = y1 - 0.035, y1 - 0.155
    rect(STX, bb, BX1, bt, fill=KEY_FILL, edge=INK, lw=0.6, r=0.03)
    text((STX + BX1) / 2, (bt + bb) / 2, [("Accuracy %.2f" % S["acc"], "b")], size=7.5, ha="center",
         box=(STX, bb, BX1, bt))
    text(STX, bb - 0.095, [("H", "i"), (" %.2f   " % S["H"], ""), ("FA %.2f" % S["FA"], "", ACC)], size=7.5)
    cs = ("%+.2f" % S["c"]).replace("-", "\u2212").replace("+0.00", "0.00")
    text(STX, bb - 0.205, [("d", "i"), ("\u2032 %.1f   " % S["dprime"], ""), ("c", "i"), (" " + cs, "")], size=7.5)


plot_scenario(A, 0.99, 1.44, "A")
plot_scenario(B, 0.33, 0.78, "B")
text((PLX0 + PLX1) / 2, 0.105, [("Evidence that the object is present \u2192", "")], size=7, ha="center", color=MUT)

# ---- (c) d' and c across stages ---------------------------------------------
text(CX0, TITLE_Y, [("(c)", "b"), (" Over stages", "")], size=8.5, name="title_c")
text(CX0, 1.690, [("Accuracy %.2f\u2013%.2f" % (min(ACCS), max(ACCS)), "")], size=7, color=MUT)
text(CX0, 1.585, [("at every stage", "")], size=7, color=MUT)
QX0, QX1, QY0, QY1 = 4.76, 5.30, 0.33, 1.44
YLO, YHI = -0.50, 0.06
qx_ = lambda s: QX0 + s / 6 * (QX1 - QX0)
qy_ = lambda v: QY0 + (v - YLO) / (YHI - YLO) * (QY1 - QY0)
line([(QX0, QY0), (QX0, QY1)], INK, 0.6)
line([(QX0, QY0), (QX1, QY0)], INK, 0.6)
for v, lab in ((0.0, "0.0"), (-0.2, "\u22120.2"), (-0.4, "\u22120.4")):
    line([(QX0 - 0.03, qy_(v)), (QX0, qy_(v))], INK, 0.6)
    text(QX0 - 0.045, qy_(v), [(lab, "")], size=7, ha="right")
line([(QX0, qy_(0)), (QX1, qy_(0))], RULE, 0.6, dash=True)
for s in range(7):
    line([(qx_(s), QY0), (qx_(s), QY0 - 0.03)], INK, 0.6)
    text(qx_(s), QY0 - 0.09, [(str(s), "")], size=7, ha="center")
text((QX0 + QX1) / 2, 0.105, [("Stage", "")], size=7, ha="center", color=MUT)
text(4.475, (QY0 + QY1) / 2, [("Change from base (SD)", "")], size=7, ha="center", rot=90, color=MUT)
for series, col in ((DD, PRES), (DC, ACC)):
    line([(qx_(s), qy_(v)) for s, v in enumerate(series)], col, 1.2)
    for s, v in enumerate(series):
        dot(qx_(s), qy_(v), col)
text(QX1 + 0.05, qy_(DC[-1]), [("c", "i", ACC)], size=8)
text(QX1 + 0.05, qy_(DD[-1]), [("d", "i", PRES), ("\u2032", "", PRES)], size=8)


# --------------------------------------------------------------------------
# matplotlib renderer + layout check
# --------------------------------------------------------------------------
def render_mpl():
    fig = plt.figure(figsize=(W, H))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W)
    ax.set_ylim(0, H)
    ax.axis("off")
    rr = fig.canvas.get_renderer()
    boxes = []

    def run_w(s, size, style):
        t = ax.text(0, 0, s.replace(" ", " "), fontsize=size, fontfamily=FONT,
                    fontstyle="italic" if "i" in style else "normal",
                    fontweight="bold" if "b" in style else "normal")
        w = t.get_window_extent(rr).width / fig.dpi
        t.remove()
        return w

    def draw_runs(it):
        ws = [run_w(s, it["size"], st) for s, st, _ in it["runs"]]
        tot = sum(ws)
        if it["rot"]:
            s = "".join(r[0] for r in it["runs"])
            t = ax.text(it["x"], it["y"], s, fontsize=it["size"], fontfamily=FONT, rotation=it["rot"],
                        ha="center", va="center", color=it["runs"][0][2], rotation_mode="anchor")
            boxes.append((t.get_window_extent(rr), it))
            return tot
        x = it["x"] - {"left": 0, "center": tot / 2, "right": tot}[it["ha"]]
        for (s, st, col), w in zip(it["runs"], ws):
            t = ax.text(x, it["y"], s.replace(" ", " "), fontsize=it["size"], fontfamily=FONT,
                        fontstyle="italic" if "i" in st else "normal",
                        fontweight="bold" if "b" in st else "normal", color=col,
                        ha="left", va="center_baseline" if it["va"] == "center" else it["va"])
            boxes.append((t.get_window_extent(rr), it))
            x += w
        return tot

    for it in P:
        k = it["k"]
        if k == "text":
            draw_runs(it)
        elif k == "rect":
            ax.add_patch(FancyBboxPatch((it["x0"], it["y0"]), it["x1"] - it["x0"], it["y1"] - it["y0"],
                                        boxstyle="round,pad=0,rounding_size=%g" % it["r"] if it["r"] else "square,pad=0",
                                        fc=it["fill"] or "none", ec=it["edge"] or "none", lw=it["lw"]))
        elif k == "line":
            xs, ys = zip(*it["pts"])
            ax.plot(xs, ys, color=it["color"], lw=it["lw"], ls=(0, (2.5, 2)) if it["dash"] else "-",
                    solid_capstyle="round", solid_joinstyle="round")
        elif k == "arrow":
            ax.annotate("", xy=it["p1"], xytext=it["p0"],
                        arrowprops=dict(arrowstyle="<|-|>" if it["both"] else "-|>", color=it["color"],
                                        lw=it["lw"], shrinkA=0, shrinkB=0, mutation_scale=5))
        elif k == "poly":
            ax.add_patch(Polygon(it["pts"], closed=True, fc=it["fill"], alpha=it["alpha"], ec="none"))
        elif k == "dot":
            ax.add_patch(matplotlib.patches.Circle((it["x"], it["y"]), it["d"] / 2, fc=it["color"], ec="white", lw=0.4,
                                                   zorder=5))
        elif k == "image":
            ax.imshow(plt.imread(it["path"]), extent=(it["x0"], it["x1"], it["y0"], it["y1"]), aspect="auto",
                      interpolation="lanczos", zorder=1)
        elif k == "fracbar":
            wn = sum(run_w(s, it["size"], st) for s, st, *_ in it["num"])
            wd = sum(run_w(s, it["size"], st) for s, st, *_ in it["den"])
            half = max(wn, wd) / 2 + 0.02
            ax.plot([it["xc"] - half, it["xc"] + half], [it["y"], it["y"]], color=INK, lw=0.6)
            bar = matplotlib.transforms.Bbox.from_extents((it["xc"] - half) * fig.dpi, (it["y"] - 0.006) * fig.dpi,
                                                          (it["xc"] + half) * fig.dpi, (it["y"] + 0.006) * fig.dpi)
            boxes.append((bar, dict(runs=[("fraction bar", "", INK)], box=(AX0, 0, AX1, H))))
        elif k == "legend":
            x, y = it["x"], it["y"]
            right = x
            for col, a, lab in it["items"]:
                ax.add_patch(FancyBboxPatch((x, y - 0.035), 0.15, 0.07, boxstyle="square,pad=0",
                                            fc=col, alpha=a, ec=col, lw=0.6))
                t = ax.text(x + 0.17, y, lab, fontsize=it["size"], fontfamily=FONT, ha="left",
                            va="center_baseline")
                bb = t.get_window_extent(rr)
                boxes.append((bb, dict(it, runs=[(lab, "", INK)])))
                right = max(right, x + 0.17 + bb.width / fig.dpi)
                y += it["dy"]
            it["_right"] = right
    ax.set_zorder(0)

    # ---- layout check: every text box inside the canvas and its container, no overlaps
    dpi = fig.dpi
    inch = [((b.x0 / dpi, b.y0 / dpi, b.x1 / dpi, b.y1 / dpi), it) for b, it in boxes]
    problems = []
    for (x0, y0, x1, y1), it in inch:
        if x0 < -0.001 or y0 < -0.001 or x1 > W + 0.001 or y1 > H + 0.001:
            problems.append("outside canvas: %r" % "".join(r[0] for r in it["runs"]))
        if it.get("box"):
            bx0, by0, bx1, by1 = it["box"]
            if x0 < bx0 + 0.02 or x1 > bx1 - 0.02:
                problems.append("text wider than its cell (%.3f in slack): %r"
                                % (min(x0 - bx0, bx1 - x1), it["runs"][0][0]))
    for i in range(len(inch)):
        for j in range(i + 1, len(inch)):
            (a0, b0, a1, b1), ia = inch[i]
            (c0, d0, c1, d1), ib = inch[j]
            if ia is ib:
                continue
            ox = min(a1, c1) - max(a0, c0)
            oy = min(b1, d1) - max(b0, d0)
            if ox > 0.004 and oy > 0.004:
                problems.append("overlap: %r / %r" % ("".join(r[0] for r in ia["runs"]),
                                                     "".join(r[0] for r in ib["runs"])))
    for it in P:
        if it["k"] == "legend" and it["_right"] > BX1 + 0.001:
            problems.append("legend overruns panel (b) by %.3f in" % (it["_right"] - BX1))
    if problems:
        raise SystemExit("LAYOUT CHECK FAILED:\n  " + "\n  ".join(problems))
    os.makedirs(OUT, exist_ok=True)
    for ext in ("pdf", "png"):
        fig.savefig(os.path.join(OUT, "fig_concept." + ext), dpi=300, facecolor="white")
    plt.close(fig)
    return len(inch)


# --------------------------------------------------------------------------
# python-pptx renderer: native, editable shapes at the same inch coordinates
# --------------------------------------------------------------------------
def render_pptx():
    from lxml import etree
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    E = lambda v: Emu(int(round(v * 914400)))
    T = lambda y: H - y
    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))

    fig = plt.figure(figsize=(W, H))
    rr = fig.canvas.get_renderer()

    def run_w(s, size, style):
        t = fig.text(0, 0, s.replace(" ", " "), fontsize=size, fontfamily=FONT,
                     fontstyle="italic" if "i" in style else "normal",
                     fontweight="bold" if "b" in style else "normal")
        w = t.get_window_extent(rr).width / fig.dpi
        t.remove()
        return w

    def set_alpha(fill_elm_parent, alpha):
        clr = fill_elm_parent.find(".//" + qn("a:srgbClr"))
        a = etree.SubElement(clr, qn("a:alpha"))
        a.set("val", str(int(alpha * 100000)))

    def add_text(x, y, runs, size, ha, rot=0):
        # Zero internal margins, so the text edge IS the box edge: anchor the box at
        # x exactly and put all the slack on the side away from the anchor.
        # (An earlier version split the slack both ways and shifted every
        # left/right-anchored string by 0.03 in.)
        w = sum(run_w(s, size, st) for s, st, _ in runs) + 0.06
        h = size * 1.25 / 72
        left = {"left": x, "center": x - w / 2, "right": x - w}[ha] if not rot else x - w / 2
        tb = sl.shapes.add_textbox(E(left), E(T(y) - h / 2), E(w), E(h))
        tf = tb.text_frame
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[ha if not rot else "center"]
        for s, st, col in runs:
            r = p.add_run()
            r.text = s
            f = r.font
            f.name, f.size, f.bold, f.italic = FONT, Pt(size), "b" in st, "i" in st
            f.color.rgb = rgb(col)
        if rot:
            tb.rotation = -rot
        return tb

    def add_line(pts, color, lw, dash=False):
        if len(pts) == 2:
            (x0, y0), (x1, y1) = pts
            c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x0), E(T(y0)), E(x1), E(T(y1)))
            c.line.color.rgb = rgb(color)
            c.line.width = Pt(lw)
            if dash:
                from pptx.enum.dml import MSO_LINE
                c.line.dash_style = MSO_LINE.DASH
            return c
        fb = sl.shapes.build_freeform(E(pts[0][0]), E(T(pts[0][1])), scale=1.0)
        fb.add_line_segments([(E(x), E(T(y))) for x, y in pts[1:]], close=False)
        s = fb.convert_to_shape()
        s.fill.background()
        s.line.color.rgb = rgb(color)
        s.line.width = Pt(lw)
        return s

    for it in P:
        k = it["k"]
        if k == "text":
            add_text(it["x"], it["y"], it["runs"], it["size"], it["ha"], it["rot"])
        elif k == "rect":
            shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if it["r"] else MSO_SHAPE.RECTANGLE,
                                      E(it["x0"]), E(T(it["y1"])), E(it["x1"] - it["x0"]), E(it["y1"] - it["y0"]))
            if it["r"]:
                shp.adjustments[0] = it["r"] / min(it["x1"] - it["x0"], it["y1"] - it["y0"])
            if it["fill"]:
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(it["fill"])
            else:
                shp.fill.background()
            if it["edge"]:
                shp.line.color.rgb = rgb(it["edge"])
                shp.line.width = Pt(it["lw"])
            else:
                shp.line.fill.background()
            shp.shadow.inherit = False
            shp.text_frame.text = ""
        elif k == "line":
            add_line(it["pts"], it["color"], it["lw"], it["dash"])
        elif k == "arrow":
            c = add_line([it["p0"], it["p1"]], it["color"], it["lw"])
            ln = c.line._get_or_add_ln()
            for tag in (["a:headEnd", "a:tailEnd"] if it["both"] else ["a:tailEnd"]):
                e = etree.SubElement(ln, qn(tag))
                e.set("type", "triangle")
                e.set("w", "sm")
                e.set("len", "sm")
        elif k == "poly":
            fb = sl.shapes.build_freeform(E(it["pts"][0][0]), E(T(it["pts"][0][1])), scale=1.0)
            fb.add_line_segments([(E(x), E(T(y))) for x, y in it["pts"][1:]], close=True)
            s = fb.convert_to_shape()
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(it["fill"])
            set_alpha(s.fill._xPr, it["alpha"])
            s.line.fill.background()
            s.shadow.inherit = False
        elif k == "dot":
            d = it["d"]
            s = sl.shapes.add_shape(MSO_SHAPE.OVAL, E(it["x"] - d / 2), E(T(it["y"]) - d / 2), E(d), E(d))
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(it["color"])
            s.line.color.rgb = rgb("#FFFFFF")
            s.line.width = Pt(0.4)
            s.shadow.inherit = False
        elif k == "image":
            sl.shapes.add_picture(it["path"], E(it["x0"]), E(T(it["y1"])), E(it["x1"] - it["x0"]),
                                  E(it["y1"] - it["y0"]))
        elif k == "fracbar":
            wn = sum(run_w(s, it["size"], st) for s, st, *_ in it["num"])
            wd = sum(run_w(s, it["size"], st) for s, st, *_ in it["den"])
            half = max(wn, wd) / 2 + 0.02
            add_line([(it["xc"] - half, it["y"]), (it["xc"] + half, it["y"])], INK, 0.6)
        elif k == "legend":
            x, y = it["x"], it["y"]
            for col, a, lab in it["items"]:
                s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(T(y + 0.035)), E(0.15), E(0.07))
                s.fill.solid()
                s.fill.fore_color.rgb = rgb(col)
                set_alpha(s.fill._xPr, a)
                s.line.color.rgb = rgb(col)
                s.line.width = Pt(0.6)
                s.shadow.inherit = False
                add_text(x + 0.17, y, [(lab, "", INK)], it["size"], "left")
                y += it["dy"]
    for shp in sl.shapes:
        st = shp._element.find(qn("p:style"))
        if st is not None:
            shp._element.remove(st)
    plt.close(fig)
    path = os.path.join(OUT, "fig_concept.pptx")
    prs.save(path)
    return path


if __name__ == "__main__":
    n = render_mpl()
    pp = render_pptx()
    print("wrote %s/fig_concept.{pdf,png} at %.2f x %.2f in  (%d text boxes, layout check passed)" % (OUT, W, H, n))
    print("wrote %s  (native, editable shapes; same coordinates)" % pp)
    print("(b) A: d'=%.1f c=%+.2f H=%.3f FA=%.3f acc=%.4f | B: d'=%.1f c=%+.2f H=%.3f FA=%.3f acc=%.4f"
          % (A["dprime"], A["c"], A["H"], A["FA"], A["acc"], B["dprime"], B["c"], B["H"], B["FA"], B["acc"]))
    print("(c) %s: accuracy %.4f-%.4f, delta-c range %.3f, delta-d' range %.3f"
          % (CELL, min(ACCS), max(ACCS), max(DC) - min(DC), max(DD) - min(DD)))
